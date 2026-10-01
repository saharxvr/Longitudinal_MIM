from __future__ import annotations

import json
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


NAVY = RGBColor(20, 44, 74)
BLUE = RGBColor(43, 108, 176)
GREEN = RGBColor(0, 145, 110)
RED = RGBColor(190, 48, 48)
DARK = RGBColor(35, 35, 35)
LIGHT = RGBColor(241, 245, 249)
WHITE = RGBColor(255, 255, 255)


def add_title(slide, title: str, subtitle: str | None = None) -> None:
    title_box = slide.shapes.add_textbox(Inches(0.55), Inches(0.25), Inches(12.2), Inches(0.55))
    paragraph = title_box.text_frame.paragraphs[0]
    paragraph.text = title
    paragraph.font.size = Pt(25)
    paragraph.font.bold = True
    paragraph.font.color.rgb = NAVY
    if subtitle:
        subtitle_box = slide.shapes.add_textbox(
            Inches(0.58), Inches(0.82), Inches(12.0), Inches(0.35)
        )
        paragraph = subtitle_box.text_frame.paragraphs[0]
        paragraph.text = subtitle
        paragraph.font.size = Pt(11)
        paragraph.font.color.rgb = DARK


def add_footer(slide, text: str = "Research comparison; not a clinical validation study") -> None:
    box = slide.shapes.add_textbox(Inches(0.5), Inches(7.18), Inches(12.3), Inches(0.2))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.size = Pt(8)
    paragraph.font.color.rgb = RGBColor(105, 105, 105)
    paragraph.alignment = PP_ALIGN.RIGHT


def add_bullets(slide, bullets: list[str], left: float, top: float, width: float, height: float) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    for index, bullet in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = bullet
        paragraph.level = 0
        paragraph.font.size = Pt(18)
        paragraph.font.color.rgb = DARK
        paragraph.space_after = Pt(11)


def add_metric_card(slide, x: float, y: float, value: str, label: str, color: RGBColor) -> None:
    shape = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(2.35), Inches(1.25))
    shape.fill.solid()
    shape.fill.fore_color.rgb = LIGHT
    shape.line.color.rgb = color
    shape.line.width = Pt(2)
    value_box = slide.shapes.add_textbox(Inches(x + 0.1), Inches(y + 0.12), Inches(2.15), Inches(0.55))
    paragraph = value_box.text_frame.paragraphs[0]
    paragraph.text = value
    paragraph.alignment = PP_ALIGN.CENTER
    paragraph.font.size = Pt(27)
    paragraph.font.bold = True
    paragraph.font.color.rgb = color
    label_box = slide.shapes.add_textbox(Inches(x + 0.12), Inches(y + 0.73), Inches(2.1), Inches(0.38))
    paragraph = label_box.text_frame.paragraphs[0]
    paragraph.text = label
    paragraph.alignment = PP_ALIGN.CENTER
    paragraph.font.size = Pt(10)
    paragraph.font.color.rgb = DARK


def create_recall_chart(summary: dict, output_path: Path) -> None:
    model = summary["consensus_sensitivity"]["ICU Model"]
    precision = summary["consensus_sensitivity"]["Claude Precision"]
    levels = np.arange(1, 6)
    fig, axes = plt.subplots(1, 2, figsize=(11, 4.1), dpi=180)
    for axis, sign, title in (
        (axes[0], "pos", "Positive / worsening changes"),
        (axes[1], "neg", "Negative / improving changes"),
    ):
        axis.plot(
            levels,
            model[f"recall_{sign}"],
            color="#2B6CB0",
            marker="o",
            linewidth=2.6,
            label="ICU model",
        )
        axis.plot(
            levels,
            precision[f"recall_{sign}"],
            color="#00916E",
            marker="o",
            linewidth=2.6,
            label="Claude Precision",
        )
        axis.set_ylim(0, 1.0)
        axis.set_xticks(levels)
        axis.set_xlabel("Pathologist consensus level")
        axis.set_ylabel("Sensitivity")
        axis.set_title(title, fontweight="bold")
        axis.grid(True, linestyle=":", alpha=0.5)
    axes[1].legend(frameon=False, loc="lower right")
    fig.tight_layout()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(output_path, bbox_inches="tight")
    plt.close(fig)


def add_case_slide(
    prs: Presentation,
    collage_path: Path,
    pair_number: int,
    headline: str,
    evidence: list[str],
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(
        slide,
        f"Case {pair_number}: {headline}",
        "Red = appearance/increase; green = disappearance/decrease; overlays normalized separately",
    )
    slide.shapes.add_picture(
        str(collage_path),
        Inches(0.35),
        Inches(1.12),
        width=Inches(9.15),
        height=Inches(5.72),
    )
    box = slide.shapes.add_shape(1, Inches(9.72), Inches(1.18), Inches(3.25), Inches(5.45))
    box.fill.solid()
    box.fill.fore_color.rgb = LIGHT
    box.line.color.rgb = RGBColor(210, 218, 228)
    add_bullets(slide, evidence, 9.93, 1.45, 2.82, 4.95)
    add_footer(slide)


def build_deck() -> Path:
    repo_root = Path(__file__).resolve().parents[2]
    experiment_root = (
        repo_root
        / "python_files"
        / "annotation tool"
        / "LLM_Pathology_Heatmaps"
        / "claude_sonnet_5_full_experiment"
    )
    presentation_root = experiment_root / "presentation"
    presentation_root.mkdir(parents=True, exist_ok=True)
    summary = json.loads(
        (experiment_root / "comparison" / "summary.json").read_text(encoding="utf-8")
    )
    recall_chart = presentation_root / "model_vs_claude_recall.png"
    create_recall_chart(summary, recall_chart)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = NAVY
    title = slide.shapes.add_textbox(Inches(0.8), Inches(1.55), Inches(11.8), Inches(1.3))
    paragraph = title.text_frame.paragraphs[0]
    paragraph.text = "ICU Semantic-Difference Model vs Claude Vision"
    paragraph.font.size = Pt(34)
    paragraph.font.bold = True
    paragraph.font.color.rgb = WHITE
    subtitle = slide.shapes.add_textbox(Inches(0.83), Inches(3.0), Inches(11.3), Inches(1.2))
    paragraph = subtitle.text_frame.paragraphs[0]
    paragraph.text = (
        "100 longitudinal CXR pairs • five pathologists • blinded Claude heatmaps\n"
        "Aggregate comparison and cases where the trained model matched consensus but Claude did not"
    )
    paragraph.font.size = Pt(20)
    paragraph.font.color.rgb = RGBColor(218, 230, 242)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Study setup", "All observers were compared in the same native image coordinate space")
    add_bullets(
        slide,
        [
            "Reference: five pathologist annotations over 100 prior/current chest X-ray pairs.",
            "ICU model: latest signed semantic-difference predictions used for the LOO analysis.",
            "Claude Precision: independent blinded image review with small, localized heatmap points.",
            "Evaluation: connected-component overlap against pathologist consensus levels 1–5.",
            "Case slides below use level-5 consensus—findings marked by all five pathologists.",
        ],
        0.85,
        1.35,
        11.7,
        4.7,
    )
    add_footer(slide)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Total results: the trained ICU model is substantially more sensitive")
    slide.shapes.add_picture(
        str(recall_chart),
        Inches(0.55),
        Inches(1.1),
        width=Inches(8.1),
        height=Inches(3.8),
    )
    add_metric_card(slide, 9.0, 1.35, "0.89", "Model level-5 positive recall", BLUE)
    add_metric_card(slide, 11.35, 1.35, "0.44", "Claude level-5 positive recall", GREEN)
    add_metric_card(slide, 9.0, 3.0, "0.77", "Model level-5 negative recall", BLUE)
    add_metric_card(slide, 11.35, 3.0, "0.38", "Claude level-5 negative recall", GREEN)
    add_bullets(
        slide,
        [
            "Model advantage persists as pathologist consensus becomes stricter.",
            "At level 5, the model detects 24/27 positive and 10/13 negative findings.",
            "Claude Precision detects 12/27 positive and 5/13 negative findings.",
        ],
        0.85,
        5.05,
        11.9,
        1.65,
    )
    add_footer(slide)

    collages = experiment_root / "collages_precision" / "per_pair"
    add_case_slide(
        prs,
        collages / "pair15_collage.png",
        15,
        "Claude predicted the opposite direction",
        [
            "Level-5 positive consensus: 2 findings.",
            "ICU model detected both (2/2).",
            "Claude Precision detected 0/2 positive findings.",
            "Claude instead produced a green improvement map.",
            "Takeaway: plausible visual explanation, wrong temporal direction.",
        ],
    )
    add_case_slide(
        prs,
        collages / "pair96_collage.png",
        96,
        "only Claude Precision missed the change",
        [
            "Level-5 positive consensus: 2 findings.",
            "ICU model detected both (2/2).",
            "Claude Ellipses and Claude Direct both reported worsening.",
            "Only Claude Precision returned no clear semantic change.",
            "Current image shows marked bilateral interval opacity.",
            "Takeaway: stricter localization improved precision but introduced a false negative.",
        ],
    )
    add_case_slide(
        prs,
        collages / "pair45_collage.png",
        45,
        "Claude missed a high-consensus improvement",
        [
            "Level-5 negative consensus: 1 finding.",
            "ICU model detected the finding (1/1).",
            "Claude Precision returned no clear change.",
            "The right-sided opacity/effusion visibly improves.",
            "Takeaway: the model is stronger for subtle longitudinal resolution.",
        ],
    )
    add_case_slide(
        prs,
        collages / "pair51_collage.png",
        51,
        "all three Claude modes were wrong",
        [
            "Level-5 negative consensus: 1 improving finding.",
            "ICU model detected the improvement (1/1).",
            "Claude Ellipses: increased pleural effusion.",
            "Claude Direct: increased consolidation.",
            "Claude Precision: new consolidation.",
            "All Claude modes assigned the wrong temporal direction.",
        ],
    )
    add_case_slide(
        prs,
        collages / "pair5_collage.png",
        5,
        "Claude modes contradict each other on direction",
        [
            "Claude Ellipses: consolidation increased.",
            "Claude Direct: pleural effusion decreased.",
            "Claude Precision: new consolidation.",
            "The same blinded pair produces opposite temporal conclusions.",
            "This indicates prompt/output-format instability rather than only localization noise.",
        ],
    )
    add_case_slide(
        prs,
        collages / "pair8_collage.png",
        8,
        "Claude modes contradict each other on pathology",
        [
            "Claude Ellipses: resolved pleural effusion.",
            "Claude Direct: resolved pneumothorax.",
            "Claude Precision: resolved effusion plus new fluid overload.",
            "Modes disagree on both diagnosis and whether worsening also occurred.",
            "The trained model output is substantially more stable across the same image pair.",
        ],
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Conclusions")
    add_bullets(
        slide,
        [
            "Claude can generate interpretable, pathology-labeled semantic heatmaps from paired CXRs.",
            "Precision prompting improved localization, but not enough to match the trained ICU model.",
            "The largest gap is recall: Claude often returns no change or assigns the wrong direction.",
            "In 13 level-5 pair/sign cases, the model detected consensus findings missed by all Claude modes.",
            "Claude modes can contradict one another on presence, direction, and pathology.",
            "The ICU model remains the stronger additional observer, especially at high consensus.",
            "Best use of the LLM today: qualitative secondary explanation—not replacement of the trained model.",
        ],
        0.9,
        1.35,
        11.55,
        4.8,
    )
    add_footer(slide, "100-pair research experiment • Generated September 2026")

    output_path = presentation_root / "ICU_Model_vs_Claude_Quick_Comparison_Corrected.pptx"
    prs.save(output_path)
    return output_path


if __name__ == "__main__":
    path = build_deck()
    print(path)
