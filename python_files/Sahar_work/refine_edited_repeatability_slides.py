from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


DARK = RGBColor(35, 35, 35)
ORANGE = RGBColor(221, 107, 32)


def slide_title(slide) -> str:
    return next(
        (
            shape.text.splitlines()[0]
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False) and shape.text.strip()
        ),
        "",
    )


def find_slide(presentation: Presentation, title_prefix: str):
    for slide in presentation.slides:
        if slide_title(slide).startswith(title_prefix):
            return slide
    raise ValueError(f"Could not find slide starting with {title_prefix!r}")


def find_text_shape(slide, text_prefix: str):
    for shape in slide.shapes:
        if (
            getattr(shape, "has_text_frame", False)
            and shape.text.strip().startswith(text_prefix)
        ):
            return shape
    raise ValueError(
        f"Could not find text starting with {text_prefix!r} on {slide_title(slide)!r}"
    )


def set_labeled_points(shape, points: list[tuple[str, str]], font_size: int) -> None:
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, (label, text) in enumerate(points):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(16)
        paragraph.font.size = Pt(font_size)
        paragraph.font.color.rgb = DARK
        label_run = paragraph.add_run()
        label_run.text = f"{label}: "
        label_run.font.bold = True
        label_run.font.size = Pt(font_size)
        label_run.font.color.rgb = ORANGE
        text_run = paragraph.add_run()
        text_run.text = text
        text_run.font.size = Pt(font_size)
        text_run.font.color.rgb = DARK


def remove_slide(presentation: Presentation, slide) -> None:
    slide_id = slide.slide_id
    slide_ids = presentation.slides._sldIdLst
    for item in list(slide_ids):
        if int(item.id) == slide_id:
            presentation.part.drop_rel(item.rId)
            slide_ids.remove(item)
            return
    raise ValueError(f"Could not remove slide {slide_id}")


def refine(source: Path, output: Path) -> Path:
    presentation = Presentation(source)

    overview = find_slide(
        presentation,
        "GPT-5.4 precision output is not deterministic",
    )
    title_shape = find_text_shape(
        overview,
        "GPT-5.4 precision output is not deterministic",
    )
    title_shape.text_frame.paragraphs[0].text = (
        "Repeatability experiment: can the LLM replace the trained model?"
    )
    subtitle_shape = find_text_shape(overview, "12 fixed pairs")
    subtitle_shape.text_frame.paragraphs[0].text = (
        "The same 12 pairs and precision prompt, repeated five times in fresh blinded GPT-5.4 sessions"
    )
    body_shape = find_text_shape(overview, "Only 1/12 pairs")
    body_shape.left = Inches(0.9)
    body_shape.top = Inches(1.55)
    body_shape.width = Inches(11.5)
    body_shape.height = Inches(3.6)
    set_labeled_points(
        body_shape,
        [
            (
                "What we did",
                "Repeated the identical semantic-heatmap task five times for each sampled pair.",
            ),
            (
                "What happened",
                "11 of 12 pairs changed pathology or direction across runs; semantic agreement was only 21%.",
            ),
            (
                "Why it matters",
                "The LLM can create plausible heatmaps, but not a stable reproducible measurement. The trained ICU model is needed.",
            ),
        ],
        font_size=22,
    )

    unstable = find_slide(presentation, "Case 47:")
    unstable_title = find_text_shape(unstable, "Case 47:")
    unstable_title.text_frame.paragraphs[0].text = (
        "Same image, five answers: why the trained model is needed"
    )
    unstable_body = find_text_shape(
        unstable,
        "Five fresh GPT-5.4 precision readings.",
    )
    unstable_body.top = Inches(1.55)
    unstable_body.height = Inches(3.3)
    set_labeled_points(
        unstable_body,
        [
            ("Same input", "The image pair and precision prompt were unchanged."),
            ("Unstable output", "Five GPT-5.4 runs produced five different semantic conclusions."),
            (
                "Takeaway",
                "A result that changes between runs cannot replace the trained semantic-difference model.",
            ),
        ],
        font_size=17,
    )

    stable = find_slide(presentation, "Case 92:")
    remove_slide(presentation, stable)

    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)
    return output


if __name__ == "__main__":
    repo_root = Path(__file__).resolve().parents[2]
    presentation_root = (
        repo_root
        / "python_files"
        / "annotation tool"
        / "LLM_Pathology_Heatmaps"
        / "claude_gpt_comparison"
        / "presentation"
    )
    source_path = (
        presentation_root
        / "ICU_Model_vs_Claude_vs_GPT_5_4_With_Repeatability_Proportional.pptx"
    )
    output_path = (
        presentation_root
        / "ICU_Model_vs_Claude_vs_GPT_5_4_Edited_Brief_Repeatability.pptx"
    )
    print(refine(source_path, output_path))
