from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


DARK = RGBColor(35, 35, 35)
ORANGE = RGBColor(221, 107, 32)


def text_shapes(slide):
    return [
        shape
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False) and shape.text.strip()
    ]


def set_text(shape, text: str) -> None:
    shape.text_frame.paragraphs[0].text = text


def set_points(
    shape,
    points: list[tuple[str, str]],
    font_size: int,
    spacing: int = 12,
) -> None:
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, (label, text) in enumerate(points):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(spacing)
        label_run = paragraph.add_run()
        label_run.text = f"{label}: "
        label_run.font.bold = True
        label_run.font.size = Pt(font_size)
        label_run.font.color.rgb = ORANGE
        text_run = paragraph.add_run()
        text_run.text = text
        text_run.font.size = Pt(font_size)
        text_run.font.color.rgb = DARK


def find_shape(slide, prefix: str):
    for shape in text_shapes(slide):
        if shape.text.strip().startswith(prefix):
            return shape
    raise ValueError(f"Could not find {prefix!r} on slide")


def reframe(source: Path, output: Path) -> Path:
    presentation = Presentation(source)
    if len(presentation.slides) != 10:
        raise ValueError(
            f"Expected the edited 10-slide deck; found {len(presentation.slides)} slides"
        )

    slide = presentation.slides[0]
    shapes = text_shapes(slide)
    set_text(shapes[0], "Why not use an LLM instead of the trained model?")
    set_text(
        shapes[1],
        "Can Claude or GPT-5.4 replace a dedicated semantic-difference model?\n"
        "Two experiments: accuracy against pathologists and repeatability across runs",
    )

    slide = presentation.slides[1]
    shapes = text_shapes(slide)
    set_text(shapes[0], "How we tested whether LLMs can replace the model")
    set_text(shapes[1], "Experiment 1: accuracy on the same 100 longitudinal chest X-ray pairs")
    body = find_shape(slide, "Reference:")
    body.height = Inches(4.9)
    set_points(
        body,
        [
            (
                "Task",
                "Claude and GPT-5.4 generated signed semantic-change heatmaps directly from prior/current images.",
            ),
            (
                "Fair comparison",
                "The LLMs and trained ICU model were evaluated in the same native coordinates against five pathologists.",
            ),
            (
                "Blinding",
                "LLMs saw no human annotations, reports, model outputs, or answers from other modes.",
            ),
            (
                "Question",
                "Can an off-the-shelf LLM detect and localize interval change as reliably as the trained model?",
            ),
        ],
        font_size=18,
    )

    slide = presentation.slides[2]
    shapes = text_shapes(slide)
    set_text(shapes[0], "Experiment 1: LLMs miss more consensus changes")
    summary = find_shape(slide, "GPT-5.4 outperformed")
    summary.top = Inches(5.05)
    summary.height = Inches(1.75)
    set_points(
        summary,
        [
            (
                "Trained model",
                "Detected 24/27 worsening and 10/13 improving level-5 consensus findings.",
            ),
            (
                "LLMs",
                "The best GPT positive result was 20/27; GPT Precision found 15/27 and Claude Precision 12/27.",
            ),
            (
                "Meaning",
                "Plausible-looking LLM heatmaps still miss substantially more high-consensus changes.",
            ),
        ],
        font_size=14,
        spacing=5,
    )

    narrative_titles = {
        3: "Failure case: LLMs reverse the direction of change",
        4: "Failure case: both LLM families agree - and are wrong",
        5: "LLMs sometimes succeed, but their modes contradict one another",
        6: "Failure case: the GPT conclusion changes with output format",
        7: "Failure case: no LLM is consistently better",
    }
    for slide_index, title in narrative_titles.items():
        set_text(text_shapes(presentation.slides[slide_index])[0], title)

    slide = presentation.slides[8]
    shapes = text_shapes(slide)
    set_text(shapes[0], "Experiment 2: LLM outputs are not reproducible")
    set_text(
        shapes[1],
        "The same 12 pairs and prompt, repeated five times in fresh blinded GPT-5.4 sessions",
    )
    body = find_shape(slide, "What we did:")
    set_points(
        body,
        [
            (
                "Test",
                "Repeat the identical heatmap task five times for each sampled pair.",
            ),
            (
                "Result",
                "11/12 pairs changed pathology or direction across runs; semantic agreement was only 21%.",
            ),
            (
                "Meaning",
                "An output that changes between identical runs is not a dependable measurement and cannot replace the trained model.",
            ),
        ],
        font_size=22,
    )

    slide = presentation.slides[9]
    shapes = text_shapes(slide)
    set_text(shapes[0], "Same input, five different answers")
    panel = find_shape(slide, "Same input:")
    set_points(
        panel,
        [
            ("Input", "The image pair and precision prompt were unchanged."),
            ("LLM output", "Five GPT-5.4 runs produced five different clinical conclusions."),
            (
                "Answer",
                "We do not use an LLM as the replacement: it is less accurate and less reproducible than the trained ICU model.",
            ),
        ],
        font_size=17,
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)
    return output


if __name__ == "__main__":
    repo_root = Path(__file__).resolve().parents[2]
    presentation_root = (
        repo_root
        / "python_files"
        / "annotation tool"
        / "LLM_Pathology_Heatmaps"
        / "claude_gpt_comparison"
        / "presentation"
    )
    source_path = (
        presentation_root
        / "ICU_Model_vs_Claude_vs_GPT_5_4_Edited_Brief_Repeatability.pptx"
    )
    output_path = presentation_root / "Why_Not_Use_LLM_Experiment.pptx"
    print(reframe(source_path, output_path))
