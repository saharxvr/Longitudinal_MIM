from __future__ import annotations

import json
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


NAVY = RGBColor(20, 44, 74)
BLUE = RGBColor(43, 108, 176)
GREEN = RGBColor(0, 145, 110)
ORANGE = RGBColor(221, 107, 32)
DARK = RGBColor(35, 35, 35)
LIGHT = RGBColor(241, 245, 249)
WHITE = RGBColor(255, 255, 255)


def add_title(slide, title: str, subtitle: str | None = None) -> None:
    box = slide.shapes.add_textbox(Inches(0.55), Inches(0.24), Inches(12.2), Inches(0.55))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = title
    paragraph.font.size = Pt(24)
    paragraph.font.bold = True
    paragraph.font.color.rgb = NAVY
    if subtitle:
        box = slide.shapes.add_textbox(Inches(0.58), Inches(0.8), Inches(12.0), Inches(0.35))
        paragraph = box.text_frame.paragraphs[0]
        paragraph.text = subtitle
        paragraph.font.size = Pt(10)
        paragraph.font.color.rgb = DARK


def add_footer(slide) -> None:
    box = slide.shapes.add_textbox(Inches(0.5), Inches(7.18), Inches(12.3), Inches(0.2))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = "Research comparison; not a clinical validation study"
    paragraph.font.size = Pt(8)
    paragraph.font.color.rgb = RGBColor(105, 105, 105)
    paragraph.alignment = PP_ALIGN.RIGHT


def add_bullets(
    slide,
    bullets: list[str],
    left: float,
    top: float,
    width: float,
    height: float,
    font_size: int = 17,
) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    for index, bullet in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = bullet
        paragraph.font.size = Pt(font_size)
        paragraph.font.color.rgb = DARK
        paragraph.space_after = Pt(10)


def create_recall_chart(claude: dict, gpt: dict, output_path: Path) -> None:
    series = {
        "ICU model": (gpt["ICU Model"], "#2B6CB0"),
        "Claude Precision": (claude["Claude Precision"], "#00916E"),
        "GPT-5.4 Precision": (gpt["GPT-5.4 Precision"], "#DD6B20"),
        "GPT-5.4 Direct": (gpt["GPT-5.4 Heatmap"], "#805AD5"),
    }
    levels = np.arange(1, 6)
    fig, axes = plt.subplots(1, 2, figsize=(12, 4.4), dpi=180)
    for axis, sign, title in (
        (axes[0], "pos", "Positive / worsening"),
        (axes[1], "neg", "Negative / improving"),
    ):
        for name, (metrics, color) in series.items():
            axis.plot(
                levels,
                metrics[f"recall_{sign}"],
                color=color,
                marker="o",
                linewidth=2.2,
                label=name,
            )
        axis.set_ylim(0, 1.0)
        axis.set_xticks(levels)
        axis.set_xlabel("Pathologist consensus level")
        axis.set_ylabel("Sensitivity")
        axis.set_title(title, fontweight="bold")
        axis.grid(True, linestyle=":", alpha=0.5)
    axes[1].legend(frameon=False, fontsize=8, loc="lower right")
    fig.tight_layout()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(output_path, bbox_inches="tight")
    plt.close(fig)


def add_metric_card(slide, x: float, y: float, value: str, label: str, color: RGBColor) -> None:
    shape = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(2.25), Inches(1.15))
    shape.fill.solid()
    shape.fill.fore_color.rgb = LIGHT
    shape.line.color.rgb = color
    value_box = slide.shapes.add_textbox(Inches(x + 0.08), Inches(y + 0.08), Inches(2.09), Inches(0.5))
    paragraph = value_box.text_frame.paragraphs[0]
    paragraph.text = value
    paragraph.alignment = PP_ALIGN.CENTER
    paragraph.font.size = Pt(25)
    paragraph.font.bold = True
    paragraph.font.color.rgb = color
    label_box = slide.shapes.add_textbox(Inches(x + 0.1), Inches(y + 0.66), Inches(2.05), Inches(0.35))
    paragraph = label_box.text_frame.paragraphs[0]
    paragraph.text = label
    paragraph.alignment = PP_ALIGN.CENTER
    paragraph.font.size = Pt(9)
    paragraph.font.color.rgb = DARK


def add_picture_contain(
    slide,
    image_path: Path,
    left: float,
    top: float,
    width: float,
    height: float,
):
    with Image.open(image_path) as image:
        image_ratio = image.width / image.height
    box_ratio = width / height
    if image_ratio >= box_ratio:
        picture_width = width
        picture_height = width / image_ratio
        picture_left = left
        picture_top = top + (height - picture_height) / 2
    else:
        picture_height = height
        picture_width = height * image_ratio
        picture_left = left + (width - picture_width) / 2
        picture_top = top
    return slide.shapes.add_picture(
        str(image_path),
        Inches(picture_left),
        Inches(picture_top),
        width=Inches(picture_width),
        height=Inches(picture_height),
    )


def add_case_slide(
    presentation: Presentation,
    collage_path: Path,
    pair_number: int,
    headline: str,
    evidence: list[str],
) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    add_title(
        slide,
        f"Case {pair_number}: {headline}",
        "All overlays normalized separately; red = worsening, green = improving",
    )
    add_picture_contain(
        slide,
        collage_path,
        0.25,
        1.08,
        9.25,
        6.05,
    )
    shape = slide.shapes.add_shape(1, Inches(9.7), Inches(1.12), Inches(3.28), Inches(5.85))
    shape.fill.solid()
    shape.fill.fore_color.rgb = LIGHT
    shape.line.color.rgb = RGBColor(210, 218, 228)
    add_bullets(slide, evidence, 9.9, 1.36, 2.88, 5.3, font_size=14)
    add_footer(slide)


def add_repeatability_slides(
    presentation: Presentation,
    repeatability_root: Path,
    repeatability_summary: dict,
) -> None:
    repeatability = repeatability_summary["primary_fresh_run_metrics"]
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    add_title(
        slide,
        "GPT-5.4 precision output is not deterministic",
        "12 fixed pairs - five fresh isolated sessions - 120 run-pair comparisons",
    )
    add_metric_card(
        slide,
        0.65,
        1.25,
        f"{repeatability['semantic_signature_agreement']:.0%}",
        "Semantic conclusion agreement",
        ORANGE,
    )
    add_metric_card(
        slide,
        3.05,
        1.25,
        f"{repeatability['any_change_agreement']:.0%}",
        "Any-change agreement",
        GREEN,
    )
    add_metric_card(
        slide,
        5.45,
        1.25,
        f"{repeatability['positive_presence_agreement']:.0%}",
        "Worsening-presence agreement",
        ORANGE,
    )
    add_metric_card(
        slide,
        7.85,
        1.25,
        f"{repeatability['negative_presence_agreement']:.0%}",
        "Improving-presence agreement",
        ORANGE,
    )
    add_metric_card(
        slide,
        10.25,
        1.25,
        f"{repeatability['signed_map_correlation']:.2f}",
        "Signed-map correlation",
        ORANGE,
    )
    add_bullets(
        slide,
        [
            "Only 1/12 pairs had the same pathology and direction in all five fresh runs.",
            "No pair produced exactly identical JSON across all five runs.",
            "Each pair generated 3.67 distinct semantic conclusions on average.",
            "The result is more stable for change/no-change than for direction, pathology, or localization.",
            "Cache mitigation: a new session and unique non-semantic nonce per run; baseline and other runs were hidden.",
            "Provider cache controls, cache-hit telemetry, temperature, and generation seed are not exposed.",
        ],
        0.85,
        3.0,
        11.7,
        3.65,
        font_size=15,
    )
    add_footer(slide)

    repeatability_collages = repeatability_root / "analysis" / "collages"
    add_case_slide(
        presentation,
        repeatability_collages / "pair47_repeatability.png",
        47,
        "five runs produce five semantic conclusions",
        [
            "Five fresh GPT-5.4 precision readings.",
            "Five distinct pathology/direction signatures.",
            "Worsening-presence agreement: 40%.",
            "Improving-presence agreement: 40%.",
            "Pathology Jaccard: 0.10.",
            "Signed-map correlation: -0.05.",
            "Variability is semantic and spatial, not merely different wording.",
        ],
    )
    add_case_slide(
        presentation,
        repeatability_collages / "pair92_repeatability.png",
        92,
        "the strongest repeatability case",
        [
            "All five runs share one semantic signature.",
            "Change and direction agreement: 100%.",
            "Pathology Jaccard: 1.00.",
            "Signed-map correlation: 0.57.",
            "Positive support Dice: 0.64.",
            "Exact JSON still differs because wording and point placement vary.",
            "Even the most stable case is not pixel-identical.",
        ],
    )


def build_deck() -> Path:
    repo_root = Path(__file__).resolve().parents[2]
    experiments_root = (
        repo_root / "python_files" / "annotation tool" / "LLM_Pathology_Heatmaps"
    )
    claude_root = experiments_root / "claude_sonnet_5_full_experiment"
    gpt_root = experiments_root / "gpt_5_4_full_experiment"
    cross_root = experiments_root / "claude_gpt_comparison"
    repeatability_root = experiments_root / "gpt_5_4_precision_repeatability"
    presentation_root = cross_root / "presentation"
    presentation_root.mkdir(parents=True, exist_ok=True)

    claude_summary = json.loads(
        (claude_root / "comparison" / "summary.json").read_text(encoding="utf-8")
    )
    gpt_summary = json.loads(
        (gpt_root / "comparison" / "summary.json").read_text(encoding="utf-8")
    )
    cross_summary = json.loads(
        (cross_root / "summary.json").read_text(encoding="utf-8")
    )
    repeatability_summary = json.loads(
        (repeatability_root / "analysis" / "summary.json").read_text(
            encoding="utf-8"
        )
    )
    chart_path = presentation_root / "model_claude_gpt_recall.png"
    create_recall_chart(
        claude_summary["consensus_sensitivity"],
        gpt_summary["consensus_sensitivity"],
        chart_path,
    )

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    box = slide.shapes.add_textbox(Inches(0.8), Inches(1.45), Inches(11.8), Inches(1.4))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = "ICU Model vs Claude Vision vs GPT-5.4"
    paragraph.font.size = Pt(34)
    paragraph.font.bold = True
    paragraph.font.color.rgb = WHITE
    box = slide.shapes.add_textbox(Inches(0.83), Inches(2.95), Inches(11.5), Inches(1.4))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = (
        "100 longitudinal chest X-ray pairs - five pathologists\n"
        "Three independent blinded output modes per vision LLM"
    )
    paragraph.font.size = Pt(20)
    paragraph.font.color.rgb = RGBColor(218, 230, 242)

    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    add_title(slide, "Study design", "Identical images, prompts, coordinate contract, and evaluation")
    add_bullets(
        slide,
        [
            "Reference: five pathologists over 100 prior/current chest X-ray pairs.",
            "Comparator: latest ICU signed semantic-difference model used for LOO statistics.",
            "Claude and GPT-5.4 each read every pair independently in ellipse, direct, and precision modes.",
            "No LLM arm saw human labels, model outputs, another LLM, or another output mode.",
            "Every heatmap was mapped back to native image coordinates before evaluation.",
            "Coordinate audit passed for 100/100 pairs for both LLMs, including non-square pairs 99-100.",
        ],
        0.85,
        1.3,
        11.7,
        5.3,
    )
    add_footer(slide)

    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    add_title(slide, "Aggregate sensitivity against pathologist consensus")
    add_picture_contain(
        slide,
        chart_path,
        0.4,
        1.05,
        8.45,
        4.05,
    )
    add_metric_card(slide, 9.0, 1.25, "0.89", "ICU level-5 positive", BLUE)
    add_metric_card(slide, 11.0, 1.25, "0.77", "ICU level-5 negative", BLUE)
    add_metric_card(slide, 9.0, 2.72, "0.56", "GPT precision positive", ORANGE)
    add_metric_card(slide, 11.0, 2.72, "0.54", "GPT precision negative", ORANGE)
    add_metric_card(slide, 9.0, 4.19, "0.44", "Claude precision positive", GREEN)
    add_metric_card(slide, 11.0, 4.19, "0.38", "Claude precision negative", GREEN)
    add_bullets(
        slide,
        [
            "GPT-5.4 outperformed Claude Precision at level 5, but remained below the ICU model.",
            "GPT direct heatmaps had the best GPT positive recall: 20/27 (0.74).",
        ],
        0.75,
        5.35,
        11.9,
        1.25,
        font_size=15,
    )
    add_footer(slide)

    agreement = cross_summary["corresponding_arm_agreement"]
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    add_title(slide, "Claude and GPT often disagree on the same blinded pair")
    add_metric_card(
        slide, 0.8, 1.35, f"{agreement['precision']['positive_presence_agreement']:.0%}",
        "Precision positive presence", ORANGE
    )
    add_metric_card(
        slide, 3.25, 1.35, f"{agreement['precision']['negative_presence_agreement']:.0%}",
        "Precision negative presence", ORANGE
    )
    add_metric_card(
        slide, 5.7, 1.35, f"{agreement['heatmap']['positive_presence_agreement']:.0%}",
        "Direct positive presence", GREEN
    )
    add_metric_card(
        slide, 8.15, 1.35, f"{agreement['heatmap']['negative_presence_agreement']:.0%}",
        "Direct negative presence", GREEN
    )
    add_bullets(
        slide,
        [
            "Agreement is measured between corresponding Claude and GPT output modes.",
            "Direct-map per-pair detection PAI appears high partly because both maps are spatially broad.",
            "Presence agreement is much lower: only 48% for positive precision-map findings.",
            "The models can disagree on change presence, temporal direction, pathology, and location.",
            "This variability supports treating vision LLM outputs as model-dependent annotations.",
        ],
        0.9,
        3.1,
        11.55,
        3.1,
    )
    add_footer(slide)

    collage_root = cross_root / "collages"
    add_case_slide(
        presentation,
        collage_root / "pair15_collage.png",
        15,
        "all LLM modes miss or reverse two level-5 worsening findings",
        [
            "Pathologists: two positive findings at level 5.",
            "ICU model: detected both (2/2).",
            "Claude Precision: fluid overload decreased.",
            "GPT Direct and Precision: fluid overload decreased.",
            "GPT Ellipses: no clear change.",
            "The principal LLM error is temporal direction, not only localization.",
        ],
    )
    add_case_slide(
        presentation,
        collage_root / "pair51_collage.png",
        51,
        "both LLM families assign the wrong direction",
        [
            "Pathologists: one level-5 improving finding.",
            "ICU model: detected it (1/1).",
            "All Claude modes report new/increased disease.",
            "All GPT modes report increased disease.",
            "No LLM mode overlaps the consensus improvement.",
            "Independent LLMs can share the same plausible but incorrect interpretation.",
        ],
    )
    add_case_slide(
        presentation,
        collage_root / "pair45_collage.png",
        45,
        "GPT detects improvement that Claude Precision misses",
        [
            "Pathologists: one level-5 improving finding.",
            "ICU model: detected it (1/1).",
            "Claude Precision: no clear change.",
            "GPT Direct and Precision: decreased pleural effusion.",
            "GPT Ellipses instead reports increased consolidation.",
            "GPT improves recall, but its own modes remain inconsistent.",
        ],
    )
    add_case_slide(
        presentation,
        collage_root / "pair96_collage.png",
        96,
        "GPT Precision succeeds where Claude Precision fails",
        [
            "Pathologists: two level-5 worsening findings.",
            "ICU model: detected both (2/2).",
            "Claude Precision: no clear change.",
            "GPT Ellipses and Precision: new fluid overload; both detect 2/2.",
            "GPT Direct: no clear change.",
            "Output format materially changes the GPT conclusion.",
        ],
    )
    add_case_slide(
        presentation,
        collage_root / "pair98_collage.png",
        98,
        "Claude detects worsening that every GPT mode misses",
        [
            "Pathologists: one level-5 positive finding.",
            "ICU model and Claude Precision detect it.",
            "All three GPT modes miss the positive finding.",
            "GPT modes instead report improving consolidation/fluid overload.",
            "This is a counterexample to a uniform GPT advantage.",
        ],
    )

    add_repeatability_slides(
        presentation,
        repeatability_root,
        repeatability_summary,
    )

    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    add_title(slide, "Conclusions")
    add_bullets(
        slide,
        [
            "GPT-5.4 is a stronger additional annotator than Claude Precision on this 100-pair set.",
            "GPT-5.4 still trails the trained ICU model at strict level-5 consensus.",
            "GPT direct heatmaps maximize positive recall, while precision prompting provides smaller support.",
            "Claude and GPT disagree frequently, so LLM annotations are not interchangeable.",
            "Both LLM families make shared high-confidence direction errors on some cases.",
            "GPT-5.4 repeat readings are also unstable: only 21% semantic agreement across fresh runs.",
            "The ICU model remains the strongest quantitative observer; LLMs are best used as complementary qualitative readers.",
        ],
        0.9,
        1.35,
        11.55,
        4.9,
    )
    add_footer(slide)

    output_path = (
        presentation_root
        / "ICU_Model_vs_Claude_vs_GPT_5_4_With_Repeatability_Proportional.pptx"
    )
    presentation.save(output_path)
    return output_path


if __name__ == "__main__":
    print(build_deck())
