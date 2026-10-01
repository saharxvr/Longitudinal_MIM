from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


NAVY = RGBColor(20, 44, 74)
BLUE = RGBColor(43, 108, 176)
GREEN = RGBColor(0, 145, 110)
ORANGE = RGBColor(221, 107, 32)
DARK = RGBColor(35, 35, 35)
LIGHT = RGBColor(241, 245, 249)
MUTED = RGBColor(105, 105, 105)


def text_shapes(slide):
    return [
        shape
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False) and shape.text.strip()
    ]


def find_shape(slide, prefix: str):
    for shape in text_shapes(slide):
        if shape.text.strip().startswith(prefix):
            return shape
    raise ValueError(f"Could not find {prefix!r}")


def delete_shape(shape) -> None:
    shape._element.getparent().remove(shape._element)


def add_text(
    slide,
    text: str,
    left: float,
    top: float,
    width: float,
    height: float,
    size: int,
    color: RGBColor = DARK,
    bold: bool = False,
    alignment: PP_ALIGN = PP_ALIGN.LEFT,
) -> None:
    box = slide.shapes.add_textbox(
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = alignment
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color


def add_compact_metric_row(
    slide,
    y: float,
    label: str,
    positive: str,
    negative: str,
    color: RGBColor,
) -> None:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(9.05),
        Inches(y),
        Inches(4.05),
        Inches(0.78),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = LIGHT
    shape.line.color.rgb = color
    shape.line.width = Pt(1.5)
    add_text(slide, label, 9.20, y + 0.18, 1.65, 0.32, 12, color, True)
    add_text(
        slide,
        positive,
        10.90,
        y + 0.13,
        0.85,
        0.40,
        17,
        color,
        True,
        PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        negative,
        11.90,
        y + 0.13,
        0.85,
        0.40,
        17,
        color,
        True,
        PP_ALIGN.CENTER,
    )


def set_labeled_points(shape, points: list[tuple[str, str]]) -> None:
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    for index, (label, text) in enumerate(points):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(5)
        label_run = paragraph.add_run()
        label_run.text = f"{label}: "
        label_run.font.bold = True
        label_run.font.size = Pt(13)
        label_run.font.color.rgb = ORANGE
        text_run = paragraph.add_run()
        text_run.text = text
        text_run.font.size = Pt(13)
        text_run.font.color.rgb = DARK


def update_deck(source: Path, output: Path) -> Path:
    presentation = Presentation(source)
    slide = presentation.slides[2]
    title = text_shapes(slide)[0]
    title.text_frame.paragraphs[0].text = (
        "Experiment 1: LLMs miss changes - and broad masks can inflate hits"
    )

    metric_values = {
        "0.89",
        "0.77",
        "0.56",
        "0.54",
        "0.44",
        "0.38",
    }
    metric_labels = {
        "ICU level-5 positive",
        "ICU level-5 negative",
        "GPT precision positive",
        "GPT precision negative",
        "Claude precision positive",
        "Claude precision negative",
    }
    for shape in list(slide.shapes):
        text = (
            shape.text.strip()
            if getattr(shape, "has_text_frame", False)
            else ""
        )
        if text in metric_values or text in metric_labels:
            delete_shape(shape)
        elif (
            shape.shape_type == 1
            and shape.left >= Inches(8.9)
            and shape.top < Inches(5.5)
        ):
            delete_shape(shape)

    add_text(
        slide,
        "Level-5 recall",
        9.05,
        1.17,
        4.05,
        0.30,
        13,
        NAVY,
        True,
        PP_ALIGN.CENTER,
    )
    add_text(slide, "Worse", 10.90, 1.45, 0.85, 0.25, 9, MUTED, True, PP_ALIGN.CENTER)
    add_text(slide, "Better", 11.90, 1.45, 0.85, 0.25, 9, MUTED, True, PP_ALIGN.CENTER)
    add_compact_metric_row(slide, 1.72, "ICU model", "0.89", "0.77", BLUE)
    add_compact_metric_row(slide, 2.72, "GPT precision", "0.56", "0.54", ORANGE)
    add_compact_metric_row(slide, 3.72, "Claude precision", "0.44", "0.38", GREEN)

    summary = find_shape(slide, "Trained model:")
    summary.top = Inches(4.88)
    summary.height = Inches(1.72)
    set_labeled_points(
        summary,
        [
            (
                "Visible area",
                "median image coverage was 4.4% for a pathologist, 9.7% for the ICU model, and 20.2% for GPT Direct.",
            ),
            (
                "Evaluation mask",
                "GPT Direct covered a median 100% of the image and GPT Precision 95.9% because Gaussian tails remain nonzero.",
            ),
            (
                "Why this matters",
                "about 82% of GPT Direct/Precision support was outside every human annotation; larger masks are more likely to overlap and count as a hit.",
            ),
        ],
    )
    add_text(
        slide,
        "Coverage over 100 pairs. Visible support: |heatmap| >= 0.10. Existing overlap evaluation uses every nonzero rendered pixel.",
        0.78,
        6.66,
        11.8,
        0.25,
        8,
        MUTED,
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)
    return output


if __name__ == "__main__":
    repo_root = Path(__file__).resolve().parents[2]
    presentation_root = (
        repo_root
        / "python_files"
        / "annotation tool"
        / "LLM_Pathology_Heatmaps"
        / "claude_gpt_comparison"
        / "presentation"
    )
    source_path = presentation_root / "Why_Not_Use_LLM_Experiment.pptx"
    output_path = presentation_root / "Why_Not_Use_LLM_With_Size_Stats.pptx"
    print(update_deck(source_path, output_path))
